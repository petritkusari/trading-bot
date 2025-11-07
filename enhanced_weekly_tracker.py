#!/usr/bin/env python3
"""
Enhanced Weekly Strategy Tracker with Detailed P&L Breakdowns
"""

import yfinance as yf
import pandas as pd
from datetime import datetime, timedelta
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference

def get_current_prices():
    """Get current market prices"""
    tickers = ['VXX', 'VIXY', 'COIN', 'TSLA', 'QQQ', 'AAPL', 'SPY']
    print("Fetching current market prices...")
    prices = {}

    for ticker in tickers:
        try:
            stock = yf.Ticker(ticker)
            data = stock.history(period='5d')
            if not data.empty:
                current_price = data['Close'].iloc[-1]
                prices[ticker] = current_price
                print(f"  {ticker}: ${current_price:.2f}")
        except Exception as e:
            print(f"  Error fetching {ticker}: {e}")

    return prices


def calculate_atm_strike(price):
    """Calculate ATM strike"""
    if price < 50:
        return round(price * 2) / 2
    elif price < 200:
        return round(price / 5) * 5
    else:
        return round(price / 10) * 10


def estimate_premium(price, strike, days_to_expiry=2, iv_pct=0.40):
    """Estimate PUT premium"""
    intrinsic = max(0, strike - price)
    time_value = price * iv_pct * (days_to_expiry / 365) ** 0.5
    total_premium = intrinsic + time_value
    return round(total_premium, 2)


def create_enhanced_tracker():
    """Create enhanced weekly tracker with detailed P&L"""

    # Get current prices
    prices = get_current_prices()

    # Portfolio setup
    capital = 50000
    allocations = {
        'VXX': 0.25,
        'COIN': 0.25,
        'TSLA': 0.25,
        'QQQ': 0.15,
    }

    # Position-level stop-losses (tiered stops from analysis!)
    stop_losses = {
        'VXX': 0.30,    # 30% for VXX
        'COIN': 0.15,   # 15% for stocks
        'TSLA': 0.15,
        'QQQ': 0.15,
    }

    # Create workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Strategy Tracker"

    # Styling
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    section_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    yellow_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    green_fill = PatternFill(start_color="C6E0B4", end_color="C6E0B4", fill_type="solid")
    red_fill = PatternFill(start_color="F4CCCC", end_color="F4CCCC", fill_type="solid")

    today = datetime.now()
    friday = today + timedelta(days=(4 - today.weekday()) % 7)

    # ============ HEADER ============
    ws['A1'] = "WEEKLY OPTIONS STRATEGY - DETAILED TRACKER"
    ws['A1'].font = Font(bold=True, size=14, color="0070C0")
    ws.merge_cells('A1:J1')
    ws['A1'].alignment = Alignment(horizontal='center')

    ws['A2'] = f"Week: {today.strftime('%Y-%m-%d')} to {friday.strftime('%Y-%m-%d')}"
    ws['A2'].font = Font(size=11)
    ws.merge_cells('A2:J2')
    ws['A2'].alignment = Alignment(horizontal='center')

    ws['A3'] = f"Initial Capital: ${capital:,.0f}"
    ws['A3'].font = Font(bold=True, size=12)

    # ============ SECTION 1: MONDAY POSITIONS ============
    row = 5
    ws[f'A{row}'] = "SECTION 1: POSITIONS ENTERED (MONDAY)"
    ws[f'A{row}'].font = Font(bold=True, size=12, color="FFFFFF")
    ws[f'A{row}'].fill = PatternFill(start_color="0070C0", end_color="0070C0", fill_type="solid")
    ws.merge_cells(f'A{row}:J{row}')

    row += 1
    headers = ['Ticker', 'Allocation %', 'Capital', 'Entry Price', 'ATM Strike',
               'Premium/Share', 'Contracts', 'Total Premium', 'Stop Loss %', 'Max Loss $']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', wrap_text=True)

    row += 1
    monday_start_row = row
    position_data = []

    for ticker, alloc_pct in allocations.items():
        actual_ticker = ticker
        if ticker == 'VXX' and ticker not in prices:
            if 'VIXY' in prices:
                actual_ticker = 'VIXY'
            else:
                continue

        if actual_ticker not in prices:
            continue

        price = prices[actual_ticker]
        allocation_capital = capital * alloc_pct
        strike = calculate_atm_strike(price)

        iv_dict = {'VXX': 0.80, 'VIXY': 0.80, 'COIN': 0.60, 'TSLA': 0.50,
                   'QQQ': 0.25, 'AAPL': 0.30, 'SPY': 0.20}
        iv = iv_dict.get(actual_ticker, 0.40)

        premium = estimate_premium(price, strike, days_to_expiry=2, iv_pct=iv)
        contracts = int(allocation_capital / (strike * 100))
        total_premium = premium * 100 * contracts
        stop_loss_pct = stop_losses.get(ticker, 0.15)
        max_loss = allocation_capital * stop_loss_pct

        position_data.append({
            'ticker': actual_ticker,
            'price': price,
            'strike': strike,
            'premium': premium,
            'contracts': contracts,
            'total_premium': total_premium,
            'row': row
        })

        ws[f'A{row}'] = actual_ticker
        ws[f'B{row}'] = f"{alloc_pct*100:.0f}%"
        ws[f'C{row}'] = allocation_capital
        ws[f'C{row}'].number_format = '$#,##0'
        ws[f'D{row}'] = price
        ws[f'D{row}'].number_format = '$#,##0.00'
        ws[f'E{row}'] = strike
        ws[f'E{row}'].number_format = '$#,##0.00'
        ws[f'F{row}'] = premium
        ws[f'F{row}'].number_format = '$#,##0.00'
        ws[f'G{row}'] = contracts
        ws[f'H{row}'] = total_premium
        ws[f'H{row}'].number_format = '$#,##0.00'
        ws[f'I{row}'] = f"{stop_loss_pct*100:.0f}%"
        ws[f'J{row}'] = max_loss
        ws[f'J{row}'].number_format = '$#,##0'

        row += 1

    monday_end_row = row - 1

    # Total premium
    ws[f'A{row}'] = "TOTAL PREMIUM TO COLLECT:"
    ws[f'A{row}'].font = Font(bold=True)
    ws[f'H{row}'] = f"=SUM(H{monday_start_row}:H{monday_end_row})"
    ws[f'H{row}'].font = Font(bold=True)
    ws[f'H{row}'].number_format = '$#,##0.00'
    ws[f'H{row}'].fill = section_fill

    # ============ SECTION 2: FRIDAY RESULTS ============
    row += 3
    ws[f'A{row}'] = "SECTION 2: FRIDAY CLOSING RESULTS"
    ws[f'A{row}'].font = Font(bold=True, size=12, color="FFFFFF")
    ws[f'A{row}'].fill = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
    ws.merge_cells(f'A{row}:J{row}')

    row += 1
    ws[f'A{row}'] = "FILL IN YELLOW CELLS ON FRIDAY"
    ws[f'A{row}'].font = Font(italic=True, color="C00000")
    ws.merge_cells(f'A{row}:J{row}')

    row += 1
    headers2 = ['Ticker', 'Friday Close', 'Strike', 'Stock Move %', 'ITM?',
                'Assigned?', 'Premium', 'Assignment Loss', 'Net P&L', 'Return %']
    for col, header in enumerate(headers2, 1):
        cell = ws.cell(row=row, column=col)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', wrap_text=True)

    row += 1
    friday_start_row = row

    for i, pos in enumerate(position_data):
        ticker = pos['ticker']
        monday_row = pos['row']

        ws[f'A{row}'] = ticker

        # Friday close - USER FILLS
        ws[f'B{row}'].fill = yellow_fill
        ws[f'B{row}'].number_format = '$#,##0.00'

        # Strike (reference from Monday)
        ws[f'C{row}'] = f"=E{monday_row}"
        ws[f'C{row}'].number_format = '$#,##0.00'

        # Stock move %
        ws[f'D{row}'] = f"=(B{row}-D{monday_row})/D{monday_row}"
        ws[f'D{row}'].number_format = '0.00%'

        # ITM?
        ws[f'E{row}'] = f'=IF(B{row}<C{row},"YES","NO")'

        # Assigned? - USER FILLS
        ws[f'F{row}'].fill = yellow_fill
        ws[f'F{row}'].value = ""

        # Premium collected (from Monday)
        ws[f'G{row}'] = f"=H{monday_row}"
        ws[f'G{row}'].number_format = '$#,##0.00'

        # Assignment loss (if assigned)
        ws[f'H{row}'] = f'=IF(F{row}="YES", (C{row}-B{row})*100*G{monday_row}, 0)'
        ws[f'H{row}'].number_format = '$#,##0.00'

        # Net P&L
        ws[f'I{row}'] = f'=G{row}-H{row}'
        ws[f'I{row}'].number_format = '$#,##0.00'

        # Return % on allocated capital
        ws[f'J{row}'] = f'=I{row}/C{monday_row}'
        ws[f'J{row}'].number_format = '0.00%'

        row += 1

    friday_end_row = row - 1

    # ============ SECTION 3: DETAILED P&L BREAKDOWN ============
    row += 2
    ws[f'A{row}'] = "SECTION 3: DETAILED P&L BREAKDOWN"
    ws[f'A{row}'].font = Font(bold=True, size=12, color="FFFFFF")
    ws[f'A{row}'].fill = PatternFill(start_color="00B050", end_color="00B050", fill_type="solid")
    ws.merge_cells(f'A{row}:E{row}')

    row += 1

    # Weekly summary
    summaries = [
        ("Total Premium Collected:", f"=SUM(G{friday_start_row}:G{friday_end_row})", green_fill),
        ("Total Assignment Losses:", f"=SUM(H{friday_start_row}:H{friday_end_row})", red_fill),
        ("NET WEEKLY P&L:", f"=SUM(I{friday_start_row}:I{friday_end_row})", green_fill),
        ("Weekly Return %:", f"=C{row+2}/{capital}", green_fill),
        ("", "", None),
        ("Annualized Return (52 weeks):", f"=POWER(1+C{row+3},52)-1", section_fill),
        ("Win Rate:", f'=COUNTIF(I{friday_start_row}:I{friday_end_row},">0")/COUNTA(I{friday_start_row}:I{friday_end_row})', section_fill),
        ("Best Position:", f'=MAX(I{friday_start_row}:I{friday_end_row})', section_fill),
        ("Worst Position:", f'=MIN(I{friday_start_row}:I{friday_end_row})', red_fill),
    ]

    for label, formula, fill in summaries:
        if label:
            ws[f'A{row}'] = label
            ws[f'A{row}'].font = Font(bold=True)
            if formula:
                ws[f'C{row}'] = formula
                ws[f'C{row}'].font = Font(bold=True)
                if fill:
                    ws[f'C{row}'].fill = fill
                if "%" in label:
                    ws[f'C{row}'].number_format = '0.00%'
                else:
                    ws[f'C{row}'].number_format = '$#,##0.00'
        row += 1

    # ============ SECTION 4: POSITION ANALYSIS ============
    row += 2
    ws[f'A{row}'] = "SECTION 4: POSITION-LEVEL ANALYSIS"
    ws[f'A{row}'].font = Font(bold=True, size=12, color="FFFFFF")
    ws[f'A{row}'].fill = PatternFill(start_color="7030A0", end_color="7030A0", fill_type="solid")
    ws.merge_cells(f'A{row}:H{row}')

    row += 1
    headers3 = ['Ticker', 'Allocated', 'P&L $', 'Return %', 'Risk Level',
                'Stop Hit?', 'Distance to Stop', 'Status']
    for col, header in enumerate(headers3, 1):
        cell = ws.cell(row=row, column=col)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', wrap_text=True)

    row += 1
    analysis_start_row = row

    for i, pos in enumerate(position_data):
        ticker = pos['ticker']
        monday_row = pos['row']
        friday_row = friday_start_row + i
        stop_loss = stop_losses.get(ticker, 0.15)

        ws[f'A{row}'] = ticker

        # Allocated capital
        ws[f'B{row}'] = f"=C{monday_row}"
        ws[f'B{row}'].number_format = '$#,##0'

        # P&L
        ws[f'C{row}'] = f"=I{friday_row}"
        ws[f'C{row}'].number_format = '$#,##0.00'

        # Return %
        ws[f'D{row}'] = f"=J{friday_row}"
        ws[f'D{row}'].number_format = '0.00%'

        # Risk level
        ws[f'E{row}'] = "VXX Hedge" if ticker in ['VXX', 'VIXY'] else "Income"

        # Stop hit?
        ws[f'F{row}'] = f'=IF(J{friday_row}<-{stop_loss},"YES","NO")'

        # Distance to stop
        ws[f'G{row}'] = f'={stop_loss}+J{friday_row}'
        ws[f'G{row}'].number_format = '0.00%'

        # Status
        ws[f'H{row}'] = f'=IF(F{row}="YES","STOPPED OUT",IF(J{friday_row}>0,"WINNER","LOSER"))'

        row += 1

    # ============ SECTION 5: INSTRUCTIONS ============
    row += 3
    ws[f'A{row}'] = "HOW TO USE THIS TRACKER:"
    ws[f'A{row}'].font = Font(bold=True, size=11, color="C00000")
    ws.merge_cells(f'A{row}:J{row}')

    row += 1
    instructions = [
        "",
        "MONDAY/TUESDAY (Setup):",
        "  1. Review Section 1 - these are the positions you would enter",
        "  2. Note the strikes, premiums, and number of contracts",
        "  3. Save this file",
        "",
        "FRIDAY (After Market Close):",
        "  OPTION A - Manual:",
        "    1. Look up closing prices for each ticker",
        "    2. Fill in YELLOW cells (Friday Close price)",
        "    3. Mark YES/NO for Assignment",
        "    4. All P&L calculates automatically!",
        "",
        "  OPTION B - Automatic (RECOMMENDED):",
        "    1. Run: python friday_price_updater.py",
        "    2. Script fetches prices and fills everything",
        "    3. Just review the results!",
        "",
        "ANALYSIS:",
        "  • Section 3 shows total P&L and key metrics",
        "  • Section 4 shows position-level analysis",
        "  • Check if any stops were hit",
        "  • Review win rate and best/worst positions",
        "",
        "NEXT WEEK:",
        "  • Run this script again for a new week",
        "  • Keep all weekly files to track progress",
        "  • After 4-6 weeks, you'll have real performance data!",
    ]

    for instruction in instructions:
        ws[f'A{row}'] = instruction
        if instruction.startswith(" "):
            ws[f'A{row}'].font = Font(size=9)
        elif ":" in instruction:
            ws[f'A{row}'].font = Font(bold=True)
        ws.merge_cells(f'A{row}:J{row}')
        row += 1

    # Set column widths
    widths = [12, 12, 12, 12, 12, 15, 10, 15, 12, 12]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    # Save
    today = datetime.now()
    filename = f"C:\\Trading\\Enhanced_Strategy_{today.strftime('%Y%m%d')}.xlsx"
    wb.save(filename)

    print(f"\n{'='*60}")
    print(f"Enhanced tracker created: {filename}")
    print(f"{'='*60}")
    print(f"\nFeatures:")
    print(f"  ✅ Detailed P&L breakdown")
    print(f"  ✅ Position-level analysis")
    print(f"  ✅ Stop-loss tracking")
    print(f"  ✅ Win rate calculation")
    print(f"  ✅ Risk metrics")
    print(f"\nOn Friday, run:")
    print(f"  python friday_price_updater.py")
    print(f"\nOr fill in yellow cells manually!")

    return filename


if __name__ == "__main__":
    create_enhanced_tracker()
