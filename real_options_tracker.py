#!/usr/bin/env python3
"""
REAL OPTIONS TRACKER - Uses actual option chain data
Fixed to show tradeable positions with real bid/ask prices
"""

import yfinance as yf
import pandas as pd
from datetime import datetime, timedelta
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

def get_real_option_data(ticker, expiry_date):
    """
    Fetch REAL option chain data for a ticker
    Returns the actual ATM PUT with real bid/ask prices
    """
    try:
        stock = yf.Ticker(ticker)

        # Get current price
        hist = stock.history(period='1d')
        if hist.empty:
            return None
        current_price = hist['Close'].iloc[-1]

        # Get option chain
        # Find the expiration date closest to our target
        expirations = stock.options
        if not expirations:
            print(f"  No options available for {ticker}")
            return None

        # Find Friday expiration
        target_exp = None
        for exp in expirations:
            exp_date = datetime.strptime(exp, '%Y-%m-%d').date()
            if exp_date >= expiry_date:
                target_exp = exp
                break

        if not target_exp:
            print(f"  No suitable expiration for {ticker}")
            return None

        # Get option chain for that expiration
        opt_chain = stock.option_chain(target_exp)
        puts = opt_chain.puts

        if puts.empty:
            return None

        # Find ATM strike (closest to current price)
        puts['distance'] = abs(puts['strike'] - current_price)
        atm_put = puts.loc[puts['distance'].idxmin()]

        return {
            'ticker': ticker,
            'current_price': current_price,
            'strike': atm_put['strike'],
            'expiry': target_exp,
            'bid': atm_put['bid'],
            'ask': atm_put['ask'],
            'last': atm_put['lastPrice'],
            'volume': atm_put['volume'],
            'openInterest': atm_put['openInterest'],
            'impliedVolatility': atm_put['impliedVolatility'],
        }

    except Exception as e:
        print(f"  Error fetching options for {ticker}: {e}")
        return None


def create_real_tracker():
    """
    Create tracker with REAL option data and tradeable positions
    """

    print("="*70)
    print("CREATING REAL OPTIONS TRACKER")
    print("="*70)

    # Portfolio setup - Using 1 contract each as user requested
    # Pick stocks that are actually tradeable
    tickers_to_try = [
        'VXX',   # Volatility - around $35
        'VIXY',  # Alternative volatility
        'AMD',   # Tech - around $140
        'PLTR',  # Tech - around $60
        'F',     # Auto - around $10
        'BAC',   # Bank - around $40
        'SOFI',  # Fintech - around $15
        'NIO',   # EV - around $5
    ]

    # Target next Friday
    today = datetime.now()
    days_until_friday = (4 - today.weekday()) % 7
    if days_until_friday == 0:
        days_until_friday = 7
    target_friday = (today + timedelta(days=days_until_friday)).date()

    print(f"\nTarget Expiration: {target_friday}")
    print("\nFetching REAL option data...")

    # Fetch real options data
    positions = []
    for ticker in tickers_to_try:
        print(f"\n  Checking {ticker}...")
        data = get_real_option_data(ticker, target_friday)
        if data:
            print(f"    Current: ${data['current_price']:.2f}")
            print(f"    Strike: ${data['strike']:.2f}")
            print(f"    Bid: ${data['bid']:.2f}  Ask: ${data['ask']:.2f}")
            print(f"    Premium (mid): ${(data['bid'] + data['ask'])/2:.2f}")
            positions.append(data)

            if len(positions) >= 4:  # Get 4 positions
                break

    if len(positions) < 4:
        print("\nWarning: Could only find", len(positions), "tradeable options")

    # Create Excel workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Real Strategy"

    # Styling
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    yellow_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    green_fill = PatternFill(start_color="C6E0B4", end_color="C6E0B4", fill_type="solid")

    # Title
    ws['A1'] = "REAL OPTIONS STRATEGY - ACTUAL MARKET DATA"
    ws['A1'].font = Font(bold=True, size=14, color="0070C0")
    ws.merge_cells('A1:J1')
    ws['A1'].alignment = Alignment(horizontal='center')

    ws['A2'] = f"Week: {today.strftime('%Y-%m-%d')} to {target_friday.strftime('%Y-%m-%d')}"
    ws['A2'].font = Font(size=11)
    ws.merge_cells('A2:J2')

    ws['A3'] = "Using 1 CONTRACT each position (as requested)"
    ws['A3'].font = Font(italic=True, color="C00000")
    ws.merge_cells('A3:J3')

    # Section 1: MONDAY POSITIONS
    row = 5
    ws[f'A{row}'] = "POSITIONS TO ENTER (ACTUAL OPTION DATA)"
    ws[f'A{row}'].font = Font(bold=True, size=12, color="FFFFFF")
    ws[f'A{row}'].fill = PatternFill(start_color="0070C0", end_color="0070C0", fill_type="solid")
    ws.merge_cells(f'A{row}:J{row}')

    row += 1
    headers = ['Ticker', 'Stock Price', 'Strike', 'Expiry', 'Bid', 'Ask',
               'Mid Price', 'Premium $', 'Contracts', 'Total Premium']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', wrap_text=True)

    row += 1
    monday_start_row = row

    # Add positions
    for pos in positions:
        mid_price = (pos['bid'] + pos['ask']) / 2
        premium_per_share = mid_price
        total_premium = premium_per_share * 100  # 1 contract = 100 shares

        ws[f'A{row}'] = pos['ticker']
        ws[f'B{row}'] = pos['current_price']
        ws[f'B{row}'].number_format = '$#,##0.00'
        ws[f'C{row}'] = pos['strike']
        ws[f'C{row}'].number_format = '$#,##0.00'
        ws[f'D{row}'] = pos['expiry']
        ws[f'E{row}'] = pos['bid']
        ws[f'E{row}'].number_format = '$#,##0.00'
        ws[f'F{row}'] = pos['ask']
        ws[f'F{row}'].number_format = '$#,##0.00'
        ws[f'G{row}'] = mid_price
        ws[f'G{row}'].number_format = '$#,##0.00'
        ws[f'H{row}'] = premium_per_share
        ws[f'H{row}'].number_format = '$#,##0.00'
        ws[f'I{row}'] = 1  # 1 contract each
        ws[f'J{row}'] = total_premium
        ws[f'J{row}'].number_format = '$#,##0.00'

        row += 1

    monday_end_row = row - 1

    # Total
    ws[f'A{row}'] = "TOTAL PREMIUM TO COLLECT:"
    ws[f'A{row}'].font = Font(bold=True)
    ws[f'J{row}'] = f"=SUM(J{monday_start_row}:J{monday_end_row})"
    ws[f'J{row}'].font = Font(bold=True)
    ws[f'J{row}'].number_format = '$#,##0.00'
    ws[f'J{row}'].fill = green_fill

    # Section 2: FRIDAY RESULTS
    row += 3
    ws[f'A{row}'] = "FRIDAY RESULTS (FILL IN YELLOW CELLS)"
    ws[f'A{row}'].font = Font(bold=True, size=12, color="FFFFFF")
    ws[f'A{row}'].fill = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
    ws.merge_cells(f'A{row}:J{row}')

    row += 1
    headers2 = ['Ticker', 'Friday Close', 'Strike', 'Move %', 'ITM?',
                'Premium Kept', 'Loss if Assigned', 'NET P&L', 'Return %', 'Status']
    for col, header in enumerate(headers2, 1):
        cell = ws.cell(row=row, column=col)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', wrap_text=True)

    row += 1
    friday_start_row = row

    for i, pos in enumerate(positions):
        monday_row = monday_start_row + i

        ws[f'A{row}'] = pos['ticker']

        # Friday close - USER FILLS
        ws[f'B{row}'].fill = yellow_fill
        ws[f'B{row}'].number_format = '$#,##0.00'
        ws[f'B{row}'].value = ""  # Empty for user to fill

        # Strike (reference)
        ws[f'C{row}'] = f"=C{monday_row}"
        ws[f'C{row}'].number_format = '$#,##0.00'

        # Move %
        ws[f'D{row}'] = f"=(B{row}-B{monday_row})/B{monday_row}"
        ws[f'D{row}'].number_format = '0.00%'

        # ITM?
        ws[f'E{row}'] = f'=IF(B{row}="","",IF(B{row}<C{row},"YES","NO"))'

        # Premium kept (always get this)
        ws[f'F{row}'] = f"=J{monday_row}"
        ws[f'F{row}'].number_format = '$#,##0.00'

        # Loss if assigned
        ws[f'G{row}'] = f'=IF(E{row}="YES",(C{row}-B{row})*100,0)'
        ws[f'G{row}'].number_format = '$#,##0.00'

        # NET P&L
        ws[f'H{row}'] = f'=F{row}-G{row}'
        ws[f'H{row}'].number_format = '$#,##0.00'

        # Return % (on capital needed for 1 contract)
        ws[f'I{row}'] = f'=H{row}/(C{row}*100)'
        ws[f'I{row}'].number_format = '0.00%'

        # Status
        ws[f'J{row}'] = f'=IF(B{row}="","PENDING",IF(H{row}>0,"WIN","LOSS"))'

        row += 1

    friday_end_row = row - 1

    # Summary
    row += 2
    ws[f'A{row}'] = "WEEKLY SUMMARY"
    ws[f'A{row}'].font = Font(bold=True, size=12, color="FFFFFF")
    ws[f'A{row}'].fill = PatternFill(start_color="00B050", end_color="00B050", fill_type="solid")
    ws.merge_cells(f'A{row}:D{row}')

    row += 1
    summaries = [
        ("Total Premium Collected:", f"=SUM(F{friday_start_row}:F{friday_end_row})", green_fill),
        ("Total Assignment Losses:", f"=SUM(G{friday_start_row}:G{friday_end_row})", None),
        ("NET WEEKLY P&L:", f"=SUM(H{friday_start_row}:H{friday_end_row})", green_fill),
        ("Average Return per Position:", f"=AVERAGE(I{friday_start_row}:I{friday_end_row})", None),
        ("Win Rate:", f'=COUNTIF(H{friday_start_row}:H{friday_end_row},">0")/COUNTA(H{friday_start_row}:H{friday_end_row})', None),
        ("Best Position:", f'=MAX(H{friday_start_row}:H{friday_end_row})', None),
        ("Worst Position:", f'=MIN(H{friday_start_row}:H{friday_end_row})', None),
    ]

    for label, formula, fill in summaries:
        ws[f'A{row}'] = label
        ws[f'A{row}'].font = Font(bold=True)
        ws[f'C{row}'] = formula
        ws[f'C{row}'].font = Font(bold=True)
        if fill:
            ws[f'C{row}'].fill = fill
        if "%" in label or "Rate" in label:
            ws[f'C{row}'].number_format = '0.00%'
        else:
            ws[f'C{row}'].number_format = '$#,##0.00'
        row += 1

    # Instructions
    row += 2
    ws[f'A{row}'] = "INSTRUCTIONS:"
    ws[f'A{row}'].font = Font(bold=True, color="C00000")
    ws.merge_cells(f'A{row}:J{row}')

    row += 1
    instructions = [
        "",
        "1. TODAY: Review the positions above (REAL option data!)",
        "2. Note: All positions are 1 CONTRACT each",
        "3. FRIDAY: Fill in the YELLOW cells (Friday closing price for each stock)",
        "4. Everything else calculates automatically!",
        "",
        "The spreadsheet will calculate:",
        "  - Whether you were assigned (stock < strike = YES)",
        "  - Premium you keep (always get this!)",
        "  - Loss if assigned (only if stock drops below strike)",
        "  - NET P&L (Premium - Loss)",
        "  - Return % on capital used",
        "  - Win/Loss status",
        "",
        "EXAMPLE:",
        "  If VXX closes at $36 (above $35 strike):",
        "    - NOT assigned",
        "    - Keep full premium (e.g., $280)",
        "    - NET P&L = +$280",
        "",
        "  If VXX closes at $32 (below $35 strike):",
        "    - ASSIGNED (buy 100 shares at $35)",
        "    - Keep premium (e.g., $280)",
        "    - Loss = ($35-$32)*100 = $300",
        "    - NET P&L = $280 - $300 = -$20",
    ]

    for instruction in instructions:
        ws[f'A{row}'] = instruction
        if instruction and instruction[0].isdigit():
            ws[f'A{row}'].font = Font(bold=True)
        ws.merge_cells(f'A{row}:J{row}')
        row += 1

    # Column widths
    widths = [10, 12, 10, 12, 8, 8, 12, 15, 12, 15]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    # Save
    filename = f"C:\\Trading\\Real_Options_Tracker_{today.strftime('%Y%m%d')}.xlsx"
    wb.save(filename)

    print(f"\n{'='*70}")
    print(f"✓ Real options tracker created: {filename}")
    print(f"{'='*70}")
    print(f"\nPositions selected ({len(positions)}):")
    total_premium = 0
    for pos in positions:
        mid = (pos['bid'] + pos['ask']) / 2
        premium = mid * 100
        total_premium += premium
        print(f"  {pos['ticker']:6s}: SELL 1 ${pos['strike']:.2f} PUT → ${premium:.0f} premium")

    print(f"\nTotal premium to collect: ${total_premium:.0f}")
    print(f"\nNext step: Fill in Friday closing prices in YELLOW cells!")

    return filename


if __name__ == "__main__":
    create_real_tracker()
