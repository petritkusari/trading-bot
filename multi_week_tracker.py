#!/usr/bin/env python3
"""
Multi-Week Cumulative Performance Tracker
Combines results from multiple weeks to show overall strategy performance
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import LineChart, Reference
from openpyxl.utils import get_column_letter
import glob
import os
from datetime import datetime

def find_completed_weeks():
    """Find all completed weekly tracker files"""
    # Look for FRIDAY_UPDATED files
    files = glob.glob("C:\\Trading\\*FRIDAY_UPDATED.xlsx")

    if not files:
        print("No completed weeks found yet!")
        print("Complete at least one week first using friday_price_updater.py")
        return []

    # Sort by creation date
    files.sort(key=os.path.getctime)

    return files


def extract_week_summary(filename):
    """Extract P&L summary from a completed week"""
    try:
        wb = openpyxl.load_workbook(filename, data_only=True)
        ws = wb.active

        # Find key metrics
        week_date = None
        total_pnl = None
        win_rate = None
        positions = []

        # Scan for data
        for row in range(1, 100):
            cell_a = ws[f'A{row}'].value
            if cell_a:
                cell_str = str(cell_a).upper()

                if "START DATE" in cell_str or "WEEK:" in cell_str:
                    week_date = ws[f'A{row}'].value

                if "WEEKLY PROFIT" in cell_str or "NET WEEKLY" in cell_str:
                    pnl_row = row
                    # Look in nearby cells for the value
                    for col in range(1, 15):
                        val = ws.cell(row=pnl_row, column=col).value
                        if val and isinstance(val, (int, float)):
                            total_pnl = val
                            break

        return {
            'filename': os.path.basename(filename),
            'date': week_date,
            'pnl': total_pnl if total_pnl else 0,
        }

    except Exception as e:
        print(f"Error reading {filename}: {e}")
        return None


def create_multi_week_summary():
    """Create cumulative multi-week performance tracker"""

    print("="*60)
    print("MULTI-WEEK CUMULATIVE TRACKER")
    print("="*60)

    # Find completed weeks
    week_files = find_completed_weeks()

    if not week_files:
        return

    print(f"\nFound {len(week_files)} completed week(s)")

    # Extract data from each week
    weeks_data = []
    for f in week_files:
        data = extract_week_summary(f)
        if data:
            weeks_data.append(data)
            print(f"  Week {len(weeks_data)}: {data['date']} - P&L: ${data['pnl']:,.0f}")

    if not weeks_data:
        print("Could not extract data from weeks!")
        return

    # Create cumulative tracker
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cumulative Performance"

    # Styling
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    green_fill = PatternFill(start_color="C6E0B4", end_color="C6E0B4", fill_type="solid")

    # Title
    ws['A1'] = "STRATEGY PERFORMANCE - CUMULATIVE RESULTS"
    ws['A1'].font = Font(bold=True, size=14, color="0070C0")
    ws.merge_cells('A1:G1')
    ws['A1'].alignment = Alignment(horizontal='center')

    ws['A2'] = f"Tracking Period: {len(weeks_data)} Week(s)"
    ws['A2'].font = Font(size=11)
    ws.merge_cells('A2:G2')

    # Weekly results table
    row = 4
    ws[f'A{row}'] = "WEEKLY RESULTS"
    ws[f'A{row}'].font = Font(bold=True, size=12)
    ws.merge_cells(f'A{row}:G{row}')

    row += 1
    headers = ['Week #', 'Date', 'Weekly P&L', 'Cumulative P&L',
               'Weekly Return %', 'Cumulative Return %', 'Capital']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', wrap_text=True)

    row += 1
    table_start_row = row
    initial_capital = 50000
    cumulative_pnl = 0
    current_capital = initial_capital

    for i, week in enumerate(weeks_data, 1):
        cumulative_pnl += week['pnl']
        weekly_return = week['pnl'] / current_capital if current_capital > 0 else 0
        cumulative_return = cumulative_pnl / initial_capital

        ws[f'A{row}'] = i
        ws[f'B{row}'] = str(week['date'])[:10] if week['date'] else f"Week {i}"
        ws[f'C{row}'] = week['pnl']
        ws[f'C{row}'].number_format = '$#,##0.00'
        ws[f'D{row}'] = cumulative_pnl
        ws[f'D{row}'].number_format = '$#,##0.00'
        ws[f'E{row}'] = weekly_return
        ws[f'E{row}'].number_format = '0.00%'
        ws[f'F{row}'] = cumulative_return
        ws[f'F{row}'].number_format = '0.00%'

        current_capital = initial_capital + cumulative_pnl
        ws[f'G{row}'] = current_capital
        ws[f'G{row}'].number_format = '$#,##0.00'

        row += 1

    table_end_row = row - 1

    # Summary statistics
    row += 2
    ws[f'A{row}'] = "PERFORMANCE SUMMARY"
    ws[f'A{row}'].font = Font(bold=True, size=12)
    ws.merge_cells(f'A{row}:D{row}')

    row += 1
    stats = [
        ("Total P&L:", f"=D{table_end_row}", green_fill),
        ("Total Return %:", f"=F{table_end_row}", green_fill),
        ("Current Capital:", f"=G{table_end_row}", green_fill),
        ("", "", None),
        ("Average Weekly P&L:", f"=AVERAGE(C{table_start_row}:C{table_end_row})", None),
        ("Average Weekly Return %:", f"=AVERAGE(E{table_start_row}:E{table_end_row})", None),
        ("Win Rate:", f'=COUNTIF(C{table_start_row}:C{table_end_row},">0")/COUNTA(C{table_start_row}:C{table_end_row})', None),
        ("Best Week:", f"=MAX(C{table_start_row}:C{table_end_row})", None),
        ("Worst Week:", f"=MIN(C{table_start_row}:C{table_end_row})", None),
        ("", "", None),
        ("Annualized Return (projected):", f"=POWER(1+F{table_end_row}/{len(weeks_data)},52)-1", green_fill),
        ("Target (backtested):", "334%", None),
        ("Difference:", f'={get_column_letter(3)}{row+10}-3.34', None),
    ]

    for label, formula, fill in stats:
        if label:
            ws[f'A{row}'] = label
            ws[f'A{row}'].font = Font(bold=True if fill else False)
            if formula:
                ws[f'C{row}'] = formula
                ws[f'C{row}'].font = Font(bold=True if fill else False)
                if fill:
                    ws[f'C{row}'].fill = fill
                if "%" in label:
                    ws[f'C{row}'].number_format = '0.00%'
                else:
                    ws[f'C{row}'].number_format = '$#,##0.00'
        row += 1

    # Analysis
    row += 2
    ws[f'A{row}'] = "ANALYSIS"
    ws[f'A{row}'].font = Font(bold=True, size=12)
    ws.merge_cells(f'A{row}:G{row}')

    row += 1
    analysis_points = [
        f"• You've completed {len(weeks_data)} week(s) of paper trading",
        f"• Total P&L: See above",
        f"• Win Rate: See above",
        "",
        "READY FOR LIVE TRADING CHECKLIST:",
        "  ☐ Completed 4+ weeks",
        "  ☐ Win rate ≥ 80%",
        "  ☐ Average weekly return ≥ 2%",
        "  ☐ Understand when to enter/exit",
        "  ☐ Comfortable with position sizing",
        "  ☐ Know how to handle assignments",
        "",
        "If ALL boxes checked → Ready for live trading with 25-50% of capital!",
    ]

    for point in analysis_points:
        ws[f'A{row}'] = point
        if point.startswith("•") or point.startswith("  ☐"):
            ws[f'A{row}'].font = Font(size=10)
        elif "READY" in point or "CHECKLIST" in point:
            ws[f'A{row}'].font = Font(bold=True, color="C00000")
        ws.merge_cells(f'A{row}:G{row}')
        row += 1

    # Add equity curve chart
    chart = LineChart()
    chart.title = "Equity Curve"
    chart.style = 13
    chart.y_axis.title = "Capital ($)"
    chart.x_axis.title = "Week"

    data = Reference(ws, min_col=7, min_row=table_start_row-1, max_row=table_end_row)
    cats = Reference(ws, min_col=1, min_row=table_start_row, max_row=table_end_row)

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)

    ws.add_chart(chart, f"I5")

    # Column widths
    for col in range(1, 8):
        ws.column_dimensions[get_column_letter(col)].width = 15

    # Save
    filename = f"C:\\Trading\\Cumulative_Performance_{datetime.now().strftime('%Y%m%d')}.xlsx"
    wb.save(filename)

    print(f"\n{'='*60}")
    print(f"Cumulative tracker created: {filename}")
    print(f"{'='*60}")
    print(f"\nThis file shows:")
    print(f"  • Week-by-week P&L")
    print(f"  • Cumulative performance")
    print(f"  • Win rate and statistics")
    print(f"  • Equity curve chart")
    print(f"  • Ready-for-live checklist")
    print(f"\nUpdate this file after completing each week!")

    return filename


if __name__ == "__main__":
    create_multi_week_summary()
